import csv
import sqlite3
from collections.abc import Generator
from pathlib import Path

import openpyxl
from rich.progress import track

from bissap.db import con, cursor, does_table_exist
from bissap.logging import get_logger

log = get_logger(__name__)

DENORMALIZE_SCRIPT = Path(__file__).parent / 'sql' / 'create_denormalized_view.sql'

COLUMNS_TYPES = {
    'MANDT': 'INT',
    'COUNTER': 'INT',
    'LOCNT': 'INT',
    'UFLAG': 'INT',
}

KNOWN_PRIMARY_KEYS = {
    'USR02': {'MANDT', 'BNAME'},
}


def clean_row(row: list[str]):
    row = row[1:-1]
    return [cell.strip(' \'"') for cell in row]


def read_csv(input: Path) -> Generator[list[str], None, None]:
    try:
        # Get the number of lines first, to track the progress
        with open(input, 'r', encoding='ISO8859') as f:
            lines = sum(1 for _l in f)

        # if it fails here, try to remove the encoding or specify UTF-8
        with open(input, 'r', encoding='ISO8859') as f:
            tablereader = track(
                csv.reader(f, delimiter='|', escapechar='\\', doublequote=False),
                total=lines,
            )

            # Parse the first row to check if it is a Sybase or a HANA dump.
            first_row = next(tablereader)
            second_row = next(tablereader)

            is_complicated = len(first_row) == 3 and first_row[0] == first_row[-1] == ''

            if not is_complicated:
                yield first_row
                yield second_row
                yield from tablereader
                return

            # Check whether there is a delimiter (e.g. -------) between the columns and the values.
            # In the first rows, there is a simple query that return itself (the first table's name)
            # This happens in some version of HANA DB
            has_delimiter = clean_row(first_row) == clean_row(second_row)

            if has_delimiter:
                yield clean_row(first_row)
                yield clean_row(next(tablereader))
            else:
                yield clean_row(next(tablereader))  # The table name, as the value
                yield clean_row(next(tablereader))  # The columns of the query
                next(tablereader)

            prev_row = None
            for row in tablereader:
                row = clean_row(row)

                # In Sybase, the table name is the result of a `SELECT 'NAME'` and is formated such that the column name is empty, then a delimiter and the result.
                # Here, we skip the delimiter of both SELECT (the table names and its values)
                if (
                    not has_delimiter
                    and len(row) == 1
                    and all(c == ' ' for c in row[0])
                ):
                    next(tablereader)
                    yield clean_row(next(tablereader))  # Table name
                    yield clean_row(next(tablereader))  # Columns
                    next(tablereader)
                elif has_delimiter and prev_row == row:
                    pass
                else:
                    yield row

                prev_row = row

    except StopIteration:
        return


def import_from_csv(input: Path):
    log.info('Importing the database. Be patient, it may take 5 minutes or more.')
    try:
        table_name = None
        rows = read_csv(input)
        for row in rows:
            if len(row) == 1:
                table_name = row[0]
                columns = next(rows)
                _handle_table_header(table_name, columns)
            else:
                if table_name is None:
                    raise ValueError('No header found before the values.')

                try:
                    _insert_line(table_name, row)
                except sqlite3.OperationalError:
                    log.warning(f'Invalid row: {row}. Skipping it.')
    finally:
        log.info('Finished importing the data.')
        con().commit()

    _denormalize_user_profiles()
    con().commit()


def _handle_table_header(table_name: str, columns: list[str]):
    """
    Create a new table in the database with the specified table name and columns.
    """
    # Use TEXT as the default type, instead of BLOB.
    columns_type = [COLUMNS_TYPES.get(column_name, 'TEXT') for column_name in columns]
    columns_decl = ', '.join(
        [
            f'{column_name} {column_type}'
            for (column_name, column_type) in zip(columns, columns_type)
        ]
    )

    if table_name in KNOWN_PRIMARY_KEYS:
        columns_decl += f', PRIMARY KEY ({", ".join(KNOWN_PRIMARY_KEYS[table_name])})'

    log.debug(f'New table: {table_name}({columns_decl})')

    _ = cursor().execute(f'CREATE TABLE IF NOT EXISTS {table_name}({columns_decl})')


def _insert_line(table_name: str, line: list[str]):
    nb_elements = len(line)
    placeholders = ','.join('?' * nb_elements)
    _ = cursor().execute(f'INSERT INTO {table_name} VALUES ({placeholders})', line)


def _denormalize_user_profiles():
    """
    Create a table `D_USER_PROFILES` containing denormalized data about user profiles.
    This table contains all and only the simple profiles of the user. It resolves the composite down to the simples.

    It also creates a view `PROFILE_AUTHS` bridging profiles and authorizations.
    """

    log.debug('Denormalizing the database.')
    _ = cursor().executescript(DENORMALIZE_SCRIPT.read_text())


def _try_denormalize_user_profiles():
    """
    If the required tables are present and it was not done before, denormalize the user profiles.
    """

    REQUIRED_TABLES = {'UST04', 'UST10S', 'UST10C', 'UST12', 'USR02'}

    if any(not does_table_exist(table) for table in REQUIRED_TABLES):
        return

    if does_table_exist('PROFILE_AUTHS'):
        return

    log.info(
        'All necessary tables are present. Denormalizing user profiles. Please be patient...'
    )
    _denormalize_user_profiles()


def import_from_xlsx(input: Path, table_name: str):
    log.info(f'Importing table {table_name}.')
    wb = openpyxl.load_workbook(input)
    ws = wb.active
    assert ws is not None, (
        'no sheets present in the workbook. verify that the file is correct.'
    )

    columns = [str(c.value) for c in ws[1]]
    _handle_table_header(table_name, columns)

    for row in ws.iter_rows(min_row=2):
        values = [str(c.value) for c in row]
        _insert_line(table_name, values)

    con().commit()

    _try_denormalize_user_profiles()


def import_from_tsv(input: Path, table_name: str):
    log.info(f'Importing table {table_name} from Tab-separated file.')

    with input.open(encoding='ISO8859') as f:
        reader = csv.reader(f, delimiter='\t')

        first_line = next(reader)
        # Handle BOM
        if first_line[0].startswith('\xef\xbb\xbf'):
            first_line[0] = first_line[0][3:]

        # If importing from SE16 / Data Browser, the file starts by the table name.
        if first_line and first_line[0] == 'Table:':
            expected_table_name = first_line[-1]
            if table_name != expected_table_name:
                log.warning(
                    f'Mismatched table names. Expected {table_name} (from filename) but got {expected_table_name} (in content)'
                )
            _ = next(reader)
            _ = next(reader)
        elif first_line and 'Dynamic List Display' in first_line[0]:
            # Sometime, the SQL Editor adds "DD.MM.YY Dynamic List Display" (where DD.MM.YY is today's date) at the start of the TSV file.
            _ = next(reader)
            _ = next(reader)
        else:
            # Else, it must be from ST04 / DBACOCKPIT
            pass

        tabs_per_columns: list[int] = []
        current_tab_count = 0
        columns: list[str] = []
        raw_columns = next(reader)
        for cell in raw_columns:
            cell = cell.lstrip(' ')
            if cell == '':
                current_tab_count += 1
            else:
                columns.append(cell)
                tabs_per_columns.append(current_tab_count)
                current_tab_count = 1
        tabs_per_columns.append(current_tab_count)

        _handle_table_header(table_name, columns)
        _ = next(reader)

        for row in reader:
            cells_to_skip = list(tabs_per_columns)

            values: list[str] = []
            current_tab_count = 0
            for cell in row:
                # Tabs are not escaped..
                if len(cells_to_skip) == 0:
                    values[-1] += '\t' + cell
                    continue

                if current_tab_count == cells_to_skip[0]:
                    cell = cell.lstrip(' ')
                    values.append(cell)
                    current_tab_count = 0
                    _ = cells_to_skip.pop(0)
                current_tab_count += 1

            # When values are empty, there are no delimiters at all. Append empty values manually
            if len(values) < len(columns):
                values += [''] * (len(columns) - len(values))

            # Some values contains tab (like USR12.VALS) and they are not escaped. As a band-aid, we concatenate excess values as the last value.
            if len(values) > len(columns):
                values[len(columns) - 1] = '\t'.join(values[len(columns) - 1 :])
                values = values[: len(columns)]

            _insert_line(table_name, values)

    con().commit()

    _try_denormalize_user_profiles()
